import os
import sys
import glob
import openpyxl
import pandas as pd
import numpy as np

# Ensure UTF-8 output
sys.stdout.reconfigure(encoding='utf-8')

ACV_DIR = r"c:\Users\Rald999\Documents\GitHub\nebulax_p3\PS3\02_Datasets\ACV"

def inspect_file(fpath, is_test=False):
    fname = os.path.basename(fpath)
    print(f"\n================================================================================")
    print(f"FILE: {fname}")
    print(f"================================================================================")
    
    wb = openpyxl.load_workbook(fpath, read_only=True)
    sheet_name = wb.sheetnames[0]
    sheet = wb[sheet_name]
    
    # Read headers
    header_row = next(sheet.iter_rows(values_only=True))
    headers = [str(h) if h is not None else "" for h in header_row]
    
    meta_cols = [c for c in headers if not c.startswith("Car ") or c == "Car model"]
    car_cols = [c for c in headers if c.startswith("Car ") and c != "Car model"]
    
    car_ids = sorted(list(set(c.split(" - ")[0].replace("Car ", "").strip() for c in car_cols)))
    
    row_count = 0
    sample_rows = []
    for r in sheet.iter_rows(values_only=True):
        row_count += 1
        if 1 < row_count <= 100:
            sample_rows.append(r)
    wb.close()
    
    print(f"Sheet name: '{sheet_name}'")
    print(f"Total Rows: {row_count - 1} | Total Columns: {len(headers)}")
    print(f"Metadata columns ({len(meta_cols)}): {meta_cols}")
    print(f"Identified Car IDs ({len(car_ids)}): {car_ids}")
    
    # Collect parameters for car 01 (or first car)
    first_car = car_ids[0]
    params = [c.split(" - ", 1)[1] for c in car_cols if c.startswith(f"Car {first_car} - ")]
    print(f"Parameters per car ({len(params)}):")
    for p in params:
        print(f"  - {p}")
        
    return {
        "fname": fname,
        "sheet_name": sheet_name,
        "rows": row_count - 1,
        "cols": len(headers),
        "car_ids": car_ids,
        "params": params,
        "headers": headers,
        "meta_cols": meta_cols
    }

def main():
    train_files = sorted(glob.glob(os.path.join(ACV_DIR, "Train", "*.xlsx")))
    test_files = sorted(glob.glob(os.path.join(ACV_DIR, "Test", "*.xlsx")))
    
    summaries = []
    for f in train_files:
        summaries.append(inspect_file(f, is_test=False))
    for f in test_files:
        summaries.append(inspect_file(f, is_test=True))
        
    print("\n\n" + "="*80)
    print("CROSS-FILE SCHEMA COMPARISON TABLE")
    print("="*80)
    for s in summaries:
        print(f"{s['fname']:<20} | Rows: {s['rows']:<6} | Cols: {s['cols']:<4} | Cars: {len(s['car_ids'])} | Params/Car: {len(s['params'])}")

if __name__ == "__main__":
    main()
