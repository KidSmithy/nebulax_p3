import os
import glob
import openpyxl
import pandas as pd
import numpy as np

ACV_DIR = r"c:\Users\Rald999\Documents\GitHub\nebulax_p3\PS3\02_Datasets\ACV"

def inspect_train_labels():
    print("=== TRAIN LABELS ===")
    labels_file = os.path.join(ACV_DIR, "Train_Labels.csv")
    df = pd.read_csv(labels_file)
    print(df.to_string(index=False))
    return df

def get_xlsx_info_fast(file_path):
    wb = openpyxl.load_workbook(file_path, read_only=True)
    sheet_name = wb.sheetnames[0]
    sheet = wb[sheet_name]
    
    # Read headers
    header_row = next(sheet.iter_rows(values_only=True))
    headers = [str(h) if h is not None else "" for h in header_row]
    
    # Count rows
    # In read_only mode, we can count rows efficiently
    row_count = 0
    first_few_rows = []
    last_row = None
    for r in sheet.iter_rows(values_only=True):
        row_count += 1
        if 1 < row_count <= 6:
            first_few_rows.append(r)
        last_row = r
    
    wb.close()
    return sheet_name, headers, row_count - 1, first_few_rows, last_row

def main():
    df_labels = inspect_train_labels()
    
    train_files = sorted(glob.glob(os.path.join(ACV_DIR, "Train", "*.xlsx")))
    test_files = sorted(glob.glob(os.path.join(ACV_DIR, "Test", "*.xlsx")))
    
    all_files = [("Train", f) for f in train_files] + [("Test", f) for f in test_files]
    
    file_summaries = {}
    
    for split, fpath in all_files:
        fname = os.path.basename(fpath)
        print(f"\n--- Inspecting {split}/{fname} ---")
        sheet_name, headers, num_rows, first_rows, last_row = get_xlsx_info_fast(fpath)
        
        # Identify metadata cols vs car cols
        meta_cols = [c for c in headers if not c.startswith("Car ")]
        car_cols = [c for c in headers if c.startswith("Car ")]
        
        # Identify car IDs
        car_ids = sorted(list(set(c.split(" - ")[0].replace("Car ", "") for c in car_cols)))
        
        # Check params per car
        params_per_car = {}
        for cid in car_ids:
            params_per_car[cid] = [c.split(" - ", 1)[1] for c in car_cols if c.startswith(f"Car {cid} - ")]
            
        print(f"File: {fname} | Sheet: {sheet_name} | Total Rows: {num_rows} | Total Cols: {len(headers)}")
        print(f"Metadata columns: {meta_cols}")
        print(f"Cars found ({len(car_ids)}): {car_ids}")
        print(f"Parameters per car (e.g. Car {car_ids[0]} count: {len(params_per_car[car_ids[0]])}):")
        for p in params_per_car[car_ids[0]][:10]:
            print(f"   - {p}")
        if len(params_per_car[car_ids[0]]) > 10:
            print(f"   ... and {len(params_per_car[car_ids[0]]) - 10} more.")
            
        # First row timestamps
        time_col_idx = None
        for idx, col_name in enumerate(headers):
            if any(term in col_name.lower() for term in ["time", "date", "时间"]):
                time_col_idx = idx
                break
        if time_col_idx is not None and first_rows:
            t_start = first_rows[0][time_col_idx]
            t_end = last_row[time_col_idx] if last_row else None
            print(f"Time span: Start = {t_start} | End = {t_end}")
            
        file_summaries[fname] = {
            "split": split,
            "rows": num_rows,
            "cols": len(headers),
            "cars": car_ids,
            "params_count_per_car": len(params_per_car[car_ids[0]]),
            "meta_cols": meta_cols,
            "params_sample": params_per_car[car_ids[0]]
        }
        
    # Compare parameter overlap
    print("\n=== PARAMETER OVERLAP ANALYSIS ===")
    common_params = None
    all_params_union = set()
    for fname, info in file_summaries.items():
        p_set = set(info["params_sample"])
        all_params_union.update(p_set)
        if common_params is None:
            common_params = p_set.copy()
        else:
            common_params = common_params.intersection(p_set)
            
    print(f"Total unique parameter names across all files: {len(all_params_union)}")
    print(f"Parameters present in ALL files ({len(common_params)}):")
    for p in sorted(list(common_params)):
        print(f"  * {p}")

if __name__ == "__main__":
    main()
